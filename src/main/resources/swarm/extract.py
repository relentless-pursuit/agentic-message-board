"""Extract a bounded uploaded/public source. No model or network access in this process."""
import io, json, sys, zipfile
from html.parser import HTMLParser

data=sys.stdin.buffer.read(2_000_001)
if len(data)>2_000_000:
    raise ValueError('Source exceeds 2 MB')
if data.startswith(b'%PDF'):
    from pypdf import PdfReader
    reader=PdfReader(io.BytesIO(data))
    if len(reader.pages)>25:
        raise ValueError('Use a document with at most 25 pages')
    text='\n'.join((page.extract_text() or '') for page in reader.pages)
elif data.startswith(b'PK'):
    archive=zipfile.ZipFile(io.BytesIO(data))
    if sum(i.file_size for i in archive.infolist())>16_000_000:
        raise ValueError('Expanded document is too large')
    from docx import Document
    doc=Document(io.BytesIO(data))
    text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
else:
    text=data.decode('utf-8-sig',errors='strict')
    if '<html' in text[:4000].lower() or '<!doctype html' in text[:4000].lower():
        class Visible(HTMLParser):
            def __init__(self):
                super().__init__(); self.skip=0; self.parts=[]
            def handle_starttag(self,tag,attrs):
                if tag in ('script','style','noscript','nav'): self.skip+=1
                elif tag in ('br','p','div','li','h1','h2','h3','section'): self.parts.append('\n')
            def handle_endtag(self,tag):
                if tag in ('script','style','noscript','nav'): self.skip=max(0,self.skip-1)
            def handle_data(self,value):
                if not self.skip: self.parts.append(value)
        parser=Visible();parser.feed(text);text=''.join(parser.parts)
text='\n'.join(line.strip() for line in text.splitlines() if line.strip()).strip()
print(json.dumps({'text':text[:14001],'truncated':len(text)>14000},ensure_ascii=True))
